import hashlib
from pathlib import Path
from dataclasses import dataclass
from app.helpers.logger import get_logger

logger = get_logger(__name__)
SUPPORTED_EXTENSIONS = {".md", ".pdf", ".docx"}


@dataclass
class RawDocument:
    filepath: str
    filename: str
    doc_type: str
    content: str
    content_hash: str


def compute_hash(content: str) -> str:
    return hashlib.sha256(content.encode("utf-8")).hexdigest()


def _parse_md(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _parse_pdf(path: Path) -> str:
    import fitz
    doc = fitz.open(str(path))
    return "\n".join(page.get_text() for page in doc)


def _parse_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def load_documents(folder_path: str) -> tuple[list[RawDocument], list[dict]]:
    root = Path(folder_path)
    if not root.exists():
        raise FileNotFoundError(f"Path does not exist: {folder_path}")

    all_files = [p for p in root.rglob("*") if p.is_file()]
    supported = [p for p in all_files if p.suffix.lower() in SUPPORTED_EXTENSIONS]
    skipped_ext = [p for p in all_files if p.suffix.lower() not in SUPPORTED_EXTENSIONS and p.suffix]

    logger.info("scanning %s: %d files found (%d supported, %d unsupported extension)",
                folder_path, len(all_files), len(supported), len(skipped_ext))
    for p in skipped_ext:
        logger.warning("unsupported extension skipped: %s", p.name)

    documents = []
    errors = []

    for path in supported:
        try:
            ext = path.suffix.lower()
            if ext == ".md":
                content = _parse_md(path)
            elif ext == ".pdf":
                content = _parse_pdf(path)
            elif ext == ".docx":
                content = _parse_docx(path)

            if not content.strip():
                errors.append({"file": path.name, "reason": "empty content after parsing"})
                logger.warning("empty content after parsing: %s", path.name)
                continue

            logger.debug("parsed: %s (%d chars)", path.name, len(content))
            documents.append(RawDocument(
                filepath=str(path),
                filename=path.name,
                doc_type=ext.lstrip("."),
                content=content,
                content_hash=compute_hash(content),
            ))
        except Exception as e:
            logger.error("parse failed: %s — %s", path.name, e, exc_info=True)
            errors.append({"file": path.name, "reason": str(e)})

    return documents, errors
